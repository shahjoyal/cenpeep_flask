"""
IS-13890 detection + ML tests.   Run from the project root:

    python -m unittest discover -s tests -v        (or:  python -m pytest tests)

Covers: the reference workbook (colours/legend/formulas), a values-only copy with
no colours, the "is this really an IS-13890 form" gate, horizontal tables, unit
rescues, the ML classifier (incl. the bare-"Oxygen" conflict), and that the field
ids agree across the backend catalog, the training data and the React config.
"""
import io
import os
import re
import sys
import unittest
import warnings

warnings.filterwarnings('ignore')
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

import openpyxl  # noqa: E402

from ml import training_data_is13890 as td  # noqa: E402
from ml.field_classifier import IS13890FieldClassifier  # noqa: E402
from routes import upload_is13890 as is_mod  # noqa: E402
from routes.upload_is13890 import FIELDS, extract_is13890  # noqa: E402

SAMPLE = os.path.join(ROOT, 'tests', 'fixtures', 'IS-13890_sample.xlsx')

EXPECTED = {
    'A': 1.9140136577858815, 'VM': 1.87711659409782, 'M': 50.5, 'FC': 45.7088697481163, 'SUC': 2.34,
    'GCV': 2248.92, 'S': 0.0, 'T': 182.83, 'Tdb': 26.0, 'Twb': 18.0, 'O2': 8.97, 'CO2': 11.7649, 'CO': 0.0651,
    'N2': 79.2, 'BL': 32000.0, 'ER': 2.0, 'Wze': 0.0, 'UBA': 84.78368970661371, 'UFA': 78.84463722397473,
}
# values of the CALCULATED rows in the sample — none of these may ever be read as an input
CALC_VALUES = [34.79827468280434, 1.1372268121790174, 2.0774746008708265, 9.573010246359935, 22.0, 16000.0,
               306.24218524574104, 124.65964814567289, 7.404647018579894, 68.31967899040421, 31.68032100959579]


def _bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet(rows, title='Sheet1'):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    return wb


class ReferenceWorkbook(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        with open(SAMPLE, 'rb') as fh:
            cls.res = extract_is13890(fh.read())

    def test_all_19_inputs_found(self):
        self.assertEqual(set(self.res['extracted']), set(EXPECTED))
        self.assertEqual(self.res['missingFields'], [])

    def test_values_exact(self):
        for fid, v in EXPECTED.items():
            self.assertAlmostEqual(self.res['extracted'][fid], v, places=9, msg=fid)

    def test_calculated_rows_never_read_as_inputs(self):
        got = list(self.res['extracted'].values())
        for v in CALC_VALUES:
            self.assertFalse(any(abs(g - v) < 1e-9 for g in got), f'calculated value {v} leaked into inputs')

    def test_detected_by_model_not_colour(self):
        d = self.res['fieldDetail']
        self.assertTrue(all(x['source'] == 'is_ml' for x in d.values()))
        self.assertTrue(all('colour' not in x for x in d.values()))
        self.assertIsNone(self.res['legend'])
        self.assertEqual(d['GCV']['cell'], "'IS - 13980 original'!E12")
        self.assertEqual(d['Wze']['cell'], "'IS - 13980 original'!E31")   # symbol/unit columns are swapped there

    def test_no_warnings(self):
        self.assertEqual(self.res['warnings'], [])


class NoColoursNoFormulas(unittest.TestCase):
    """The same sheet saved as plain values (fills + formulas stripped) must still give the same inputs."""

    def test_values_only_copy(self):
        wb = openpyxl.load_workbook(SAMPLE, data_only=True)
        for ws in wb:
            for row in ws.iter_rows():
                for c in row:
                    c.fill = openpyxl.styles.PatternFill(fill_type=None)
        res = extract_is13890(_bytes(wb))
        self.assertEqual(set(res['extracted']), set(EXPECTED))
        for fid, v in EXPECTED.items():
            self.assertAlmostEqual(res['extracted'][fid], v, places=9, msg=fid)
        
    def test_colours_are_irrelevant(self):
        """Same sheet with every fill painted ONE colour (and the legend gone): identical result."""
        wb = openpyxl.load_workbook(SAMPLE)
        ws = wb.worksheets[0]
        vals = openpyxl.load_workbook(SAMPLE, data_only=True).worksheets[0]
        for row in ws.iter_rows():
            for c in row:
                c.fill = openpyxl.styles.PatternFill('solid', fgColor='FFFF00')
                if isinstance(c.value, str) and c.value.startswith('='):
                    c.value = vals[c.coordinate].value          # formulas -> plain numbers
        res = extract_is13890(_bytes(wb))
        self.assertEqual(set(res['extracted']), set(EXPECTED))
        got = list(res['extracted'].values())
        for v in CALC_VALUES:       # calculated rows are rejected by the model's OUT_OF_SCOPE class, not by colour
            self.assertFalse(any(abs(g - v) < 1e-9 for g in got), v)


class Gate(unittest.TestCase):
    def test_cenpeep_like_sheet_is_not_an_is13890_form(self):
        rows = [['Parameter', 'Symbol', 'UoM', 'Value'],
                ['Ash', 'A', '%', 40], ['Moisture', 'M', '%', 12.2], ['Volatile matter', 'VM', '%', 22.9],
                ['Fixed carbon', 'FC', '%', 24.9], ['Gross calorific value', 'GCV', 'kcal/kg', 3320],
                ['O2 at APH inlet', 'O2in', '%', 3.5], ['Flue gas outlet temp', 'Tgo', 'C', 135]]
        res = extract_is13890(_bytes(_sheet(rows)))
        self.assertEqual(res['extracted'], {})
        self.assertEqual(len(res['missingFields']), len(is_mod.REQUIRED_IDS))

    def test_standard_named_in_title_is_enough(self):
        rows = [['Boiler efficiency by heat loss method (IS 13890)'], ['Parameter', 'Value'],
                ['Ash', 3], ['Moisture', 50], ['Volatile matter', 20], ['Fixed carbon', 26], ['Gross calorific value', 2200]]
        res = extract_is13890(_bytes(_sheet(rows)))
        self.assertEqual(set(res['extracted']), {'A', 'M', 'VM', 'FC', 'GCV'})

    def test_empty_workbook(self):
        res = extract_is13890(_bytes(_sheet([['nothing', 'here']])))
        self.assertEqual(res['extracted'], {})


class Layouts(unittest.TestCase):
    def test_horizontal_table_averages_readings(self):
        rows = [['Date', 'Ash %', 'Moisture %', 'Wet bulb temp', 'Dry bulb temp', 'O2 at APH inlet %', 'Evaporation ratio'],
                ['d1', 2.0, 50.0, 18, 26, 8.0, 2.0], ['d2', 4.0, 52.0, 20, 28, 10.0, 2.2]]
        res = extract_is13890(_bytes(_sheet(rows)))
        e = res['extracted']
        self.assertAlmostEqual(e['A'], 3.0)
        self.assertAlmostEqual(e['M'], 51.0)
        self.assertAlmostEqual(e['Twb'], 19.0)
        self.assertAlmostEqual(e['O2'], 9.0)
        self.assertIn('mean of 2 readings', res['fieldDetail']['A']['note'])

    def test_units_rescued_and_flagged(self):
        rows = [['Particulars', 'UoM', 'Value'],
                ['CO at APH inlet', 'ppm', 650], ['Boiler Load', 'TPH', 32], ['Evaporation ratio', 'kg/kg', 2],
                ['Wet bulb temp', 'C', 18], ['Dry bulb temp', 'C', 26], ['Ash', '%', 2], ['Moisture', '%', 50]]
        res = extract_is13890(_bytes(_sheet(rows)))
        self.assertAlmostEqual(res['extracted']['CO'], 0.065)
        self.assertAlmostEqual(res['extracted']['BL'], 32000.0)
        self.assertTrue(any('ppm' in w for w in res['warnings']))

    def test_impossible_value_is_reported_not_filled(self):
        rows = [['Particulars', 'UoM', 'Value'], ['O2 at APH inlet', '%', 55], ['Wet bulb temp', 'C', 18],
                ['Dry bulb temp', 'C', 26], ['Ash', '%', 2], ['Moisture', '%', 50], ['Evaporation ratio', '', 2]]
        res = extract_is13890(_bytes(_sheet(rows)))
        self.assertNotIn('O2', res['extracted'])
        self.assertTrue(any('outside the plausible range' in w for w in res['warnings']))

    def test_model_picks_label_not_unit_cell(self):
        rows = [['Particulars', 'Symbol', 'UoM', 'Value'], ['Loss due to unburned hydrogen', 'Luh', 'kcal/kg', 0],
                ['Wet bulb temp', 'Twb', 'C', 18], ['Dry bulb temp', 'Tdb', 'C', 26], ['Ash', 'A', '%', 2],
                ['Moisture', 'M', '%', 50], ['Evaporation ratio', '', 'kg/kg', 2]]
        res = extract_is13890(_bytes(_sheet(rows)))
        self.assertNotIn('GCV', res['extracted'])
        self.assertEqual(res['warnings'], [])

    def test_formula_is_calculation_but_pure_sheet_link_is_input(self):
        L = is_mod._LINK_RE
        self.assertTrue(L.match("='Lab Data'!C14"))
        self.assertTrue(L.match('=Lab!$B$2'))
        self.assertFalse(L.match('=E10+0.9*(E8-14)'))
        self.assertFalse(L.match("='Lab'!C14*2"))
        self.assertFalse(is_mod._is_input_cell(is_mod.Cell(5.0, '=E10+E11')))
        self.assertTrue(is_mod._is_input_cell(is_mod.Cell(5.0, "='Lab'!C14")))
        self.assertTrue(is_mod._is_input_cell(is_mod.Cell(5.0, None)))


class Labels(unittest.TestCase):
    def test_exclusions(self):
        m = is_mod._match_label
        self.assertEqual(m('Ash')[0], 'A')
        self.assertNotEqual((m('Bottom ash quantity') or ('',))[0], 'A')
        self.assertNotEqual((m('Loss due to formation of CO') or ('',))[0], 'CO')
        self.assertEqual(m('CO at APH inlet')[0], 'CO')
        self.assertEqual(m('CO2 at APH inlet')[0], 'CO2')
        self.assertEqual(m('Avg Flue Gas O/L Temp')[0], 'T')
        self.assertNotEqual((m('Flue gas inlet temp') or ('',))[0], 'T')
        self.assertEqual(m('Unburnts in bottom ash')[0], 'UBA')
        self.assertEqual(m('Unburnt in Fly ash')[0], 'UFA')
        self.assertIsNone(m('Oxygen'))                      # bare element names are calculated rows
        self.assertIsNone(m('Nitrogen'))
        self.assertIsNone(m('Nitrogen by difference'))
        self.assertEqual(m('Oxygen in flue gas')[0], 'O2')
        self.assertIsNone(m('Ambient Temp (Dry bulb + Wet bulb)/2'))


class MLClassifier(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.clf = IS13890FieldClassifier()
        cls.clf.train()

    def pred(self, text):
        return self.clf.predict(text, threshold=0.60)[0]

    def test_no_label_conflicts_in_training_data(self):
        self.assertEqual(td.label_conflicts(), {})

    def test_bare_oxygen_family_is_rejected(self):
        for t in ['Oxygen', 'Nitrogen', 'Carbon', 'Hydrogen', 'Oxygen %', 'Oxygen (%)', 'Ultimate analysis oxygen',
                  'Oxygen by difference', 'Nitrogen by difference']:
            self.assertIsNone(self.pred(t), t)

    def test_gas_classes_with_context_or_formula(self):
        for t, f in [('O2', 'O2'), ('O₂ at APH inlet', 'O2'), ('Oxygen in flue gas', 'O2'), ('Oxygen at APH inlet', 'O2'),
                     ('CO2 at APH inlet', 'CO2'), ('Carbon dioxide in flue gas', 'CO2'), ('CO', 'CO'),
                     ('Carbon monoxide at APH inlet', 'CO'), ('CO ppm', 'CO'), ('N2 at APH inlet', 'N2'),
                     ('N2 by difference', 'N2'), ('Nitrogen at APH inlet', 'N2')]:
            self.assertEqual(self.pred(t), f, t)

    def test_other_fields(self):
        for t, f in [('Boiler Air Pre Heater out let Temp', 'T'), ('Flue gas exit temperature', 'T'),
                     ('Ambient Temp (Wet Bulb)', 'Twb'), ('Ambient Temp (Dry bulb)', 'Tdb'),
                     ('Quantity of atomizing steam', 'Wze'), ('Evaporation Ratio', 'ER'), ('Boiler Load kg/h', 'BL'),
                     ('Unburnts in bottom ash', 'UBA'), ('LOI fly ash', 'UFA'), ('Sucrose', 'SUC'),
                     ('Gross Calorific Value', 'GCV'), ('Sulphur', 'S'), ('Volatile matter', 'VM'),
                     ('Fixed Carbon', 'FC'), ('Moisture', 'M'), ('Ash', 'A')]:
            self.assertEqual(self.pred(t), f, t)

    def test_calculated_rows_rejected(self):
        for t in ['Loss due to dry flue gas', 'Loss due to formation of CO', 'Total Loss %', 'Boiler Efficiency',
                  'Weight of dry refuse collected', 'HHV of total dry refuse', 'Bagasse Consumption',
                  'Mean Sp. Heat of dry Flue Gas', 'Ambient Temp (Dry bulb + Wet bulb)/2',
                  'kg of carbon burnt per kg of AF fuel', 'Vapour Enthalpy at 0.07kgf/cm2 at APH o/l',
                  'Flue gas inlet temp', 'Primary air temp', 'Steam pressure', 'Feed water temp']:
            self.assertIsNone(self.pred(t), t)


class CatalogsAgree(unittest.TestCase):
    def test_backend_and_training_ids_match(self):
        self.assertEqual(set(FIELDS), set(td.get_field_ids()))

    def test_react_config_ids_match(self):
        path = os.path.join(ROOT, 'src', 'config', 'fields.is13890.js')
        if not os.path.exists(path):
            self.skipTest('React config not present')
        src = open(path, encoding='utf-8').read()
        ids = set(re.findall(r'\bf\(\s*[\'"]([A-Za-z0-9]+)[\'"]', src))      # f('A', 'Ash', ...) field rows
        self.assertEqual(ids, set(FIELDS))


if __name__ == '__main__':
    unittest.main()
