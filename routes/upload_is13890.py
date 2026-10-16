# ════════════════════════════════════════════════════════════════════
# routes/upload_is13890.py — IS-13890 (boiler efficiency, heat-loss method)
# field detection
#
# Like routes/upload_bs2885.py this is a plain, self-contained module — NOT its
# own Flask route. Its one public function, extract_is13890(file_bytes), is
# called from routes/upload.py's /api/upload handler and adds extra
# `…IS13890` keys to that response. It shares no state with the CENPEEP / BEE /
# BS-2885 detection, so none of that can be affected by this file.
#
# What the IS-13890 workbook looks like
# -------------------------------------
#   Particulars | Symbol | UoM | Value | Formula / note
#   (In the reference workbook inputs are orange and calculations green — the
#   detector deliberately does NOT use colours; the model decides from the labels.)
#
# Detection strategy — the trained Python model decides which field a row is:
#
#   1. MODEL     every label cell is scored by the dedicated IS-13890 text
#                classifier (ml/field_classifier.py + ml/training_data_is13890.py,
#                TF-IDF char n-grams, with an OUT_OF_SCOPE class for the sheet's
#                calculated rows). Cell COLOURS are NOT used at all.
#   2. VETO      a model answer is dropped if the label contains a word that rules
#                the field out ("Ash" never becomes "Bottom ash quantity", "CO"
#                never "Loss due to formation of CO").
#   3. RULES     only when the model has no confident answer, token rules /
#                symbol+unit matching are tried as a safety net.
#   4. FORMULA   a cell holding a formula is a calculation (skipped) — unless it
#                is a bare link to another sheet (how people feed lab data in).
#   5. SANITY    physical-range and unit checks (CO in ppm, load in t/h, …); an
#                impossible value is reported as a warning, not silently filled.
#   6. GATE      a sheet is only accepted as an IS-13890 form when it carries
#                IS-13890-specific fields (wet-bulb temp, evaporation ratio,
#                unburnts in ash, atomizing steam …) or names the standard.
#
# Two layouts are understood: the vertical "Particulars | Value" form above,
# and a horizontal table (field names across a header row, readings below —
# multiple readings are averaged, like the CENPEEP parser does).
# ════════════════════════════════════════════════════════════════════
import io
import re
from collections import namedtuple

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

STANDARD_LABEL = 'IS-13890'

# ── Field catalog ───────────────────────────────────────────────────────────
# Keep ids in sync with src/config/fields.is13890.js and ml/training_data_is13890.py
# (tests/test_is13890_extract.py checks this).
#
# aliases : (phrase, weight)  — contiguous token sequence in the normalised label
# bags    : (all_tokens, any_tokens|None, weight) — tokens in ANY order
# exclude : phrases that veto the field when present in the label
# symbols : symbol-column values that identify the field (fallback)
# sym_units: unit words that must accompany a symbol-only match (None = any)
# range   : (lo, hi) physically plausible values AFTER unit conversion
FIELDS = {
    'A': dict(label='Ash', unit='%', required=True, range=(0, 60),
              aliases=[('ash', .95)],
              exclude=['bottom', 'fly', 'flyash', 'bed', 'coarse', 'pit', 'dust', 'riddling', 'loss', 'unburnt', 'refuse',
                       'weight', 'kg', 'hhv', 'softening', 'temp', 'quantity', 'collected', 'esp', 'stack', 'loi'],
              symbols=['a'], sym_units={'%'}),
    'VM': dict(label='Volatile Matter', unit='%', required=True, range=(0, 100),
               aliases=[('volatile matter', .97), ('volatiles', .9), ('vol matter', .93), ('v m', .9), ('vm', .9),
                        ('volatile', .85)],
               exclude=['loss'], symbols=['vm'], sym_units={'%'}),
    'M': dict(label='Moisture', unit='%', required=True, range=(0, 90),
              aliases=[('total moisture', .97), ('moisture', .95), ('tm', .8)],
              exclude=['loss', 'kg', 'enthalpy', 'vapour', 'hydrogen', 'per', 'air', 'humidity', 'steam', 'ambient'],
              symbols=['m', 'tm'], sym_units={'%'}),
    'FC': dict(label='Fixed Carbon', unit='%', required=True, range=(0, 100),
               aliases=[('fixed carbon', .97), ('fixed c', .9), ('fc', .9)],
               exclude=['loss'], symbols=['fc'], sym_units={'%'}),
    'SUC': dict(label='Sucrose', unit='%', required=False, range=(0, 100),
                aliases=[('sucrose', .95), ('sugar', .85), ('pol', .85)],
                exclude=['loss'], symbols=[], sym_units=None),
    'GCV': dict(label='Gross Calorific Value', unit='kcal/kg', required=True, range=(500, 12000),
                aliases=[('gross calorific value', .97), ('gross cv', .95), ('gcv', .97), ('higher heating value', .95),
                         ('higher heat value', .95), ('gross heat value', .95), ('higher calorific value', .95),
                         ('hhv', .9), ('calorific value', .8), ('heating value', .8)],
                exclude=['net', 'ncv', 'refuse', 'loss', 'radiation', 'ash'],
                symbols=['gcv', 'hhv'], sym_units=None),
    'S': dict(label='Sulphur', unit='%', required=True, range=(0, 10),
              aliases=[('sulphur', .97), ('sulfur', .97), ('salphur', .9), ('sulpher', .9)],
              exclude=['dioxide', 'so2', 'loss'], symbols=['s'], sym_units={'%'}),
    'T': dict(label='Avg Flue Gas Outlet Temp', unit='°C', required=True, range=(50, 700),
              aliases=[('flue gas outlet temp', .97), ('flue gas exit temp', .95), ('exit gas temp', .9),
                       ('stack temp', .85), ('chimney temp', .85), ('flue gas temp', .8), ('fg temp', .8),
                       ('aph outlet temp', .75), ('air heater outlet temp', .75), ('air pre heater outlet temp', .75)],
              bags=[({'gas', 'leaving', 'temp'}, None, .85)],
              exclude=['inlet', 'ambient', 'entering', 'steam', 'primary', 'secondary', 'hot', 'furnace', 'bed', 'main',
                       'feed', 'water', 'bulb'],
              symbols=['t', 'tg'], sym_units={'oc', 'c', 'deg c', 'deg', 'degc'}),
    'Tdb': dict(label='Ambient Temp (Dry bulb)', unit='°C', required=True, range=(-20, 60),
                aliases=[('ambient temp dry bulb', .97), ('dry bulb temp', .95), ('dry bulb', .93), ('dbt', .9),
                         ('db temp', .85), ('atmospheric temp', .7), ('ambient temp', .7)],
                bags=[({'air', 'entering', 'temp'}, None, .8)],
                exclude=['wet', 'wb', 'wbt', 'mean', 'average', 'avg', 'outlet', 'hot', 'primary', 'secondary', 'flue',
                         'gas', 'steam', 'feed', 'water', 'fuel'],
                symbols=[], sym_units=None),
    'Twb': dict(label='Ambient Temp (Wet Bulb)', unit='°C', required=True, range=(-20, 50),
                aliases=[('ambient temp wet bulb', .97), ('wet bulb temp', .95), ('wet bulb', .93), ('wbt', .9),
                         ('wb temp', .85)],
                exclude=['dry', 'db', 'mean', 'average', 'avg'], symbols=[], sym_units=None),
    'O2': dict(label='O₂ at APH Inlet', unit='%', required=True, range=(0, 21),
               aliases=[('o2 at aph inlet', .97), ('o2 aph inlet', .95), ('o2 at air heater inlet', .95),
                        ('oxygen at aph inlet', .95), ('o2 in flue gas', .9), ('flue gas o2', .9),
                        ('o2 at aph outlet', .7), ('o2', .8)],
               bags=[({'oxygen', 'flue'}, None, .9), ({'oxygen', 'aph'}, None, .9), ({'oxygen', 'gas'}, None, .85),
                     ({'oxygen', 'air', 'heater'}, None, .9), ({'oxygen', 'excess'}, None, .85)],
               # NB: a BARE "Oxygen" is deliberately NOT an alias — in this worksheet it is the calculated
               # ultimate-analysis row (see ml/training_data_is13890.py "bare Oxygen" note).
               exclude=['loss', 'ambient', 'dissolved', 'atomizing', 'atomising'],
               symbols=['o2'], sym_units={'%'}),
    'CO2': dict(label='CO₂ at APH Inlet', unit='%', required=True, range=(0, 25),
                aliases=[('co2 at aph inlet', .97), ('co2 aph inlet', .95), ('co2 at air heater inlet', .95),
                         ('co2 in flue gas', .9), ('flue gas co2', .9), ('co2', .85)],
                bags=[({'carbon', 'dioxide'}, {'flue', 'gas', 'aph', 'heater', 'inlet'}, .9)],
                exclude=['loss', 'so2'], symbols=['co2'], sym_units={'%'}),
    'CO': dict(label='CO at APH Inlet', unit='%', required=True, range=(0, 5),
               aliases=[('co at aph inlet', .97), ('co aph inlet', .95), ('co at air heater inlet', .95),
                        ('co in flue gas', .9), ('flue gas co', .9), ('carbon monoxide', .9), ('co', .8)],
               exclude=['loss', 'formation', 'due', 'co2'], symbols=['co'], sym_units={'%'}),
    'N2': dict(label='N₂ at APH Inlet', unit='%', required=True, range=(50, 95),
               aliases=[('n2 at aph inlet', .97), ('n2 by difference', .95), ('n2 by diff', .95), ('n2 balance', .9),
                        ('n2 at air heater inlet', .95), ('n2 in flue gas', .9), ('flue gas n2', .9), ('n2', .8),
                        ('nitrogen at aph inlet', .95), ('nitrogen in flue gas', .9)],
               # "nitrogen by difference" is NOT matched on its own: the ultimate-analysis block uses bare
               # element names, so it needs a flue-gas / APH word next to it.
               bags=[({'nitrogen', 'flue'}, None, .9), ({'nitrogen', 'aph'}, None, .9)],
               exclude=['loss'], symbols=['n2'], sym_units={'%'}),
    'BL': dict(label='Boiler Load', unit='kg/h', required=True, range=(500, 5_000_000),
               aliases=[('boiler load', .97), ('main steam flow', .9), ('boiler steam flow', .9), ('steam generation', .9),
                        ('steam generated', .85), ('steam production', .85), ('steam output', .85), ('boiler output', .85),
                        ('steam flow', .85)],
               exclude=['pressure', 'temp', 'enthalpy', 'loss', 'atomizing', 'atomising', 'feed', 'water', 'spray',
                        'bagasse', 'fuel', 'ratio'],
               symbols=[], sym_units=None),
    'ER': dict(label='Evaporation Ratio', unit='kg/kg', required=True, range=(0.3, 20),
               aliases=[('evaporation ratio', .97), ('evap ratio', .95), ('evaporation factor', .85),
                        ('evaporation rate', .8), ('steam fuel ratio', .9), ('steam to fuel ratio', .9),
                        ('steam bagasse ratio', .9), ('steam to bagasse ratio', .9)],
               bags=[({'steam', 'ratio'}, None, .8)],
               exclude=['loss'], symbols=[], sym_units=None),
    'Wze': dict(label='Quantity of Atomizing Steam', unit='kg/h', required=True, range=(0, 1_000_000),
                aliases=[('atomizing steam', .97), ('atomising steam', .97), ('atomizer steam', .95),
                         ('steam for atomizing', .95)],
                bags=[({'atomizing'}, None, .9), ({'atomising'}, None, .9), ({'atomizer'}, None, .85)],
                exclude=['loss'], symbols=['wze'], sym_units=None),
    'UBA': dict(label='Unburnts in Bottom Ash', unit='%', required=True, range=(0, 100),
                aliases=[('unburnt in bottom ash', .97), ('unburnt carbon in bottom ash', .97), ('unburnt bottom ash', .95),
                         ('loi bottom ash', .97), ('bottom ash loi', .97), ('bottom ash unburnt', .95),
                         ('carbon in bottom ash', .9), ('uba', .85)],
                bags=[({'bottom', 'ash'}, {'unburnt', 'loi', 'ignition', 'combustible', 'combustibles', 'carbon', 'ubc'}, .93),
                      ({'bed', 'ash'}, {'unburnt', 'loi', 'ignition', 'combustible', 'combustibles', 'carbon', 'ubc'}, .88)],
                exclude=['loss', 'quantity', 'temp'], symbols=['uba'], sym_units=None),
    'UFA': dict(label='Unburnts in Fly Ash', unit='%', required=True, range=(0, 100),
                aliases=[('unburnt in fly ash', .97), ('unburnt carbon in fly ash', .97), ('unburnt fly ash', .95),
                         ('loi fly ash', .97), ('fly ash loi', .97), ('fly ash unburnt', .95),
                         ('carbon in fly ash', .9), ('ufa', .85)],
                bags=[({'fly', 'ash'}, {'unburnt', 'loi', 'ignition', 'combustible', 'combustibles', 'carbon', 'ubc'}, .93),
                      ({'flyash'}, {'unburnt', 'loi', 'ignition', 'combustible', 'combustibles', 'carbon', 'ubc'}, .9)],
                exclude=['loss', 'quantity', 'temp'], symbols=['ufa'], sym_units=None),
}
FIELD_ORDER = list(FIELDS)
REQUIRED_IDS = [fid for fid in FIELD_ORDER if FIELDS[fid]['required']]

# Fields that only an IS-13890 form has (a CENPEEP / BEE log has none of them,
# or at most one by coincidence). A sheet must carry >= MIN_DISTINCTIVE_FIELDS of
# these — or name the standard in its title — to be accepted. See "GATE" above.
DISTINCTIVE_IDS = {'Tdb', 'Twb', 'ER', 'UBA', 'UFA', 'Wze', 'SUC'}
MIN_DISTINCTIVE_FIELDS = 2
STANDARD_MARKER_RE = re.compile(r'13890|13980|is\s*-?\s*13890', re.I)   # (the sample sheet's own title says "13980")
MARKER_SCAN_ROWS = 12

VERTICAL_MAX_ROWS = 300        # a worksheet form is ~65 rows; 300 is generous
VERTICAL_MAX_COLS = 16
COLUMN_HEADER_SCAN_ROWS = 40
COLUMN_MAX_ROWS = 2000
COLUMN_MAX_COLS = 60
MIN_FIELDS_TO_ACCEPT_SHEET = 4
ML_THRESHOLD_FLOOR = 0.60      # never accept an ML guess below this, whatever the classifier default is
SYMBOL_ONLY_CONFIDENCE = 0.72

Candidate = namedtuple('Candidate', 'fid value conf source row cell label unit_note')


# ── Text helpers ─────────────────────────────────────────────────────────────
def _norm(text):
    s = str(text).lower()
    s = s.replace('₂', '2').replace('°', ' ').replace('º', ' ')
    s = re.sub(r'unburn(?:t|ed|ts)\b', 'unburnt', s)
    s = re.sub(r'\btemperature\b', 'temp', s)
    s = re.sub(r'[^a-z0-9]+', ' ', s).strip()
    # o/l, i/l, "out let", "in let" -> outlet / inlet
    s = re.sub(r'\bo l\b', 'outlet', s)
    s = re.sub(r'\bi l\b', 'inlet', s)
    s = re.sub(r'\bout let\b', 'outlet', s)
    s = re.sub(r'\bin let\b', 'inlet', s)
    return s


def _tokens(norm):
    return norm.split()


def _contains_seq(tokens, seq):
    n, m = len(tokens), len(seq)
    if m == 0 or m > n:
        return False
    return any(tokens[i:i + m] == seq for i in range(n - m + 1))


# pre-tokenise the catalog once
for _fid, _spec in FIELDS.items():
    _spec['_aliases'] = [(_tokens(_norm(p)), w) for p, w in _spec['aliases']]
    _spec['_exclude'] = [_tokens(_norm(p)) for p in _spec['exclude']]
    _spec['_bags'] = [(set(a), set(b) if b else None, w) for a, b, w in _spec.get('bags', [])]


def _match_label(text):
    """Best (fid, weight, ntokens) for one label string by rules, or None."""
    norm = _norm(text)
    if not norm or len(norm) > 160:
        return None
    toks = _tokens(norm)
    tokset = set(toks)
    best = None
    for fid, spec in FIELDS.items():
        if any(_contains_seq(toks, ex) for ex in spec['_exclude']):
            continue
        w_best, n_best = 0.0, 0
        for seq, w in spec['_aliases']:
            if _contains_seq(toks, seq) and (w, len(seq)) > (w_best, n_best):
                w_best, n_best = w, len(seq)
        for all_set, any_set, w in spec['_bags']:
            if all_set <= tokset and (any_set is None or (any_set & tokset)) and (w, len(all_set)) > (w_best, n_best):
                w_best, n_best = w, len(all_set)
        if w_best and (best is None or (w_best, n_best) > (best[1], best[2])):
            best = (fid, w_best, n_best)
    return best


def _match_symbol(sym_text, unit_texts):
    sym = _norm(sym_text)
    if not sym or len(sym) > 6:
        return None
    for fid, spec in FIELDS.items():
        if sym in spec['symbols']:
            need = spec.get('sym_units')
            if need is not None:
                units = {_norm(u) for u in unit_texts} | {str(u).strip().lower() for u in unit_texts}
                if not (units & need):
                    continue
            return fid
    return None


def _to_num(val):
    if val is None or isinstance(val, bool):
        return None
    if isinstance(val, (int, float)):
        return float(val) if val == val and abs(val) != float('inf') else None
    if isinstance(val, str):
        s = val.strip().replace(',', '')
        if s.endswith('%'):
            s = s[:-1].strip()
        if not s or s.startswith('='):
            return None
        try:
            f = float(s)
            return f if f == f and abs(f) != float('inf') else None
        except ValueError:
            return None
    return None


# ── Workbook access (openpyxl, with a values-only calamine fallback for .xls) ─
# formula = the cell's formula text ("=E10+0.9*(E8-14)") or None
Cell = namedtuple('Cell', 'v formula')

# A formula that is nothing but a link to one cell of another sheet: ='Lab'!C14
_LINK_RE = re.compile(r"^=\s*(?:'[^']+'|[A-Za-z0-9_.\[\] ]+)!\$?[A-Za-z]{1,3}\$?\d+\s*$")


class _Book:
    """Uniform read access: rows_values() is cheap (no styles); rows_full()
    adds the formula text for the cells we actually inspect."""

    def __init__(self, data):
        self.kind = None
        self._wbv = self._wbf = self._cal = None
        self._data = data
        if openpyxl is not None:
            try:
                self._wbv = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
                self.kind = 'openpyxl'
                self.names = list(self._wbv.sheetnames)
            except Exception:
                self._wbv = None
        if self.kind is None:
            import python_calamine
            self._cal = python_calamine.CalamineWorkbook.from_filelike(io.BytesIO(data))
            self.kind = 'calamine'
            self.names = list(self._cal.sheet_names)

    def _worksheet(self, wb, name):
        ws = wb[name]
        return ws if hasattr(ws, 'iter_rows') else None      # chart sheets have no cell grid

    def rows_values(self, name, max_row, max_col):
        if self.kind == 'openpyxl':
            ws = self._worksheet(self._wbv, name)
            if ws is None:
                return []
            return [list(r) for r in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col, values_only=True)]
        rows = self._cal.get_sheet_by_name(name).to_python(skip_empty_area=False)[:max_row]
        return [list(r[:max_col]) for r in rows]

    def rows_full(self, name, max_row, max_col):
        """Grid of Cell(value, formula_text) — values from the cached-values pass,
        formulas from a second (formulas) pass over the same cells."""
        if self.kind == 'calamine':
            return [[Cell(v, None) for v in r] for r in self.rows_values(name, max_row, max_col)]
        if self._wbf is None:
            self._wbf = openpyxl.load_workbook(io.BytesIO(self._data), data_only=False, read_only=True)
        ws_v, ws_f = self._worksheet(self._wbv, name), self._worksheet(self._wbf, name)
        if ws_v is None or ws_f is None:
            return []
        grid = []
        for row_v, row_f in zip(ws_v.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col),
                                ws_f.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col)):
            out = []
            for cv, cf in zip(row_v, row_f):
                fv = getattr(cf, 'value', None)
                formula = fv if isinstance(fv, str) and fv.startswith('=') and fv[1:2] != '=' else None
                out.append(Cell(getattr(cv, 'value', None), formula))
            grid.append(out)
        return grid


def _col_letter(idx):
    s = ''
    idx += 1
    while idx:
        idx, rem = divmod(idx - 1, 26)
        s = chr(65 + rem) + s
    return s


# ── Unit handling ────────────────────────────────────────────────────────────
_TPH_RE = re.compile(r'(?<![a-z])(tph|t/h|t/hr|tonnes?/h|tons?/h|tonnes? per h|tons? per h|t per h)', re.I)


def _unit_adjust(fid, value, raw_texts):
    """Apply obvious unit conversions. Returns (value, note|None)."""
    joined = ' '.join(str(t) for t in raw_texts if t)
    note = None
    if fid == 'CO' and re.search(r'ppm', joined, re.I) and value > FIELDS['CO']['range'][1]:
        return value / 10000.0, 'ppm → %'
    if fid in ('BL', 'Wze') and _TPH_RE.search(joined) and not re.search(r'kg', joined, re.I):
        return value * 1000.0, 't/h → kg/h'
    return value, note


def _sane(fid, value):
    lo, hi = FIELDS[fid]['range']
    return lo <= value <= hi


def _finalise_value(fid, value, raw_texts, where, warnings):
    """Unit conversion + range check. Returns (value, note) or None if rejected."""
    value, note = _unit_adjust(fid, value, raw_texts)
    if not _sane(fid, value):
        # implicit-unit rescues, flagged to the user rather than silent
        if fid == 'CO' and 5 < value <= 50000:
            return value / 10000.0, 'looks like ppm → converted to %'
        if fid == 'BL' and 0 < value < FIELDS['BL']['range'][0] and _sane('BL', value * 1000.0):
            return value * 1000.0, 'looks like t/h → converted to kg/h'
        lo, hi = FIELDS[fid]['range']
        warnings.append(f"{FIELDS[fid]['label']}: value {value:g} at {where} is outside the plausible range "
                        f"{lo:g}–{hi:g} {FIELDS[fid]['unit']} — ignored.")
        return None
    return value, note


# ── Vertical form layout ─────────────────────────────────────────────────────
def _row_value_cell(row):
    """(col, cell) of the first numeric cell that has some label text before it."""
    seen_text = False
    for c, cell in enumerate(row):
        v = cell.v
        if isinstance(v, str) and len(v.strip()) >= 1 and _to_num(v) is None and not v.strip().startswith('='):
            seen_text = True
            continue
        if seen_text and _to_num(v) is not None and not isinstance(v, bool):
            return c, cell
    return None, None


def _is_input_cell(cell):
    """A formula is a calculation — unless it is just a link to another sheet."""
    return not (cell.formula and not _LINK_RE.match(cell.formula))


def _vetoed(fid, text):
    """True if the label contains a word that rules this field out."""
    toks = _tokens(_norm(text))
    return any(_contains_seq(toks, ex) for ex in FIELDS[fid]['_exclude'])


def _get_classifier(warnings):
    try:
        from ml.field_classifier import get_is13890_classifier, IS13890_CONFIDENCE_THRESHOLD
        return get_is13890_classifier(), max(IS13890_CONFIDENCE_THRESHOLD, ML_THRESHOLD_FLOOR)
    except Exception as e:
        msg = f'AI field detection unavailable ({type(e).__name__}); only rule-based detection was used.'
        if msg not in warnings:
            warnings.append(msg)
        return None, None


_UNIT_ONLY_RE = re.compile(r'^(%|kcal\s*/\s*kg.*|kg\s*/\s*(h|hr|kg)|t\s*/\s*h(r)?|tph|ppm|mg\s*/\s*nm3|°?\s*c|o\s*c|deg\.?\s*c|kcal)$', re.I)


def _model_pick(texts, clf, thr):
    """Best (fid, score, text) among several label texts, or None. Unit-only cells ("%", "kcal/kg") are
    never offered to the model; vetoed answers are dropped; on equal scores the longer (more descriptive)
    label wins over a bare symbol."""
    cand = [str(t) for t in texts if len(_norm(t)) >= 3 and not _UNIT_ONLY_RE.match(str(t).strip())]
    if not cand or clf is None:
        return None
    best = None
    for text, (fid, score, _m) in zip(cand, clf.predict_batch(cand, threshold=thr)):
        if not fid or _vetoed(fid, text):
            continue
        key = (round(score, 2), len(text.strip()))
        if best is None or key > best[0]:
            best = (key, fid, score, text)
    return (best[1], best[2], best[3]) if best else None


def _extract_vertical(book, name, warnings):
    """Returns {fid: Candidate} for the vertical "Particulars | Value" layout."""
    vals = book.rows_values(name, VERTICAL_MAX_ROWS, VERTICAL_MAX_COLS)
    if not vals:
        return {}
    # Cheap pre-check: does any row of this sheet even read like an IS-13890 form?
    hits = 0
    for row in vals:
        if any(isinstance(v, str) and _match_label(v) for v in row):
            hits += 1
            if hits >= 3:
                break
    if hits < 3:
        return {}
    # Guard: a wide hourly data dump is not a form
    numeric_cols_seen = max((sum(1 for c in row if isinstance(c, (int, float)) and not isinstance(c, bool))
                             for row in vals[:30]), default=0)
    if numeric_cols_seen > 8:
        return {}

    grid = book.rows_full(name, VERTICAL_MAX_ROWS, VERTICAL_MAX_COLS)
    clf, thr = _get_classifier(warnings)

    cands = []
    for r, row in enumerate(grid):
        vc, vcell = _row_value_cell(row)
        if vc is None or not _is_input_cell(vcell):
            continue
        texts = [(c, cell.v) for c, cell in enumerate(row[:vc]) if isinstance(cell.v, str) and cell.v.strip()]
        if not texts:
            continue
        raw_texts = [t for _, t in texts]
        where = f"'{name}'!{_col_letter(vc)}{r + 1}"

        fid = conf = label_text = source = None
        # 1) the trained model
        pick = _model_pick(raw_texts, clf, thr)
        if pick:
            fid, conf, label_text = pick
            source = 'is_ml'
        else:
            # 2) safety net: token rules, then symbol + unit
            best = None
            for _c, t in texts:
                m = _match_label(t)
                if m and (best is None or (m[1], m[2]) > (best[1][1], best[1][2])):
                    best = (t, m)
            if best:
                label_text, (fid, conf, _n) = best[0], best[1]
                source = 'is_label'
            else:
                for _c, t in texts:
                    f2 = _match_symbol(t, [x for x in raw_texts if x is not t])
                    if f2:
                        fid, conf, label_text, source = f2, SYMBOL_ONLY_CONFIDENCE, raw_texts[0], 'is_label'
                        break
        if not fid:
            continue
        fin = _finalise_value(fid, float(_to_num(vcell.v)), raw_texts, where, warnings)
        if fin is None:
            continue
        value, note = fin
        cands.append(Candidate(fid, value, round(min(conf, 0.99), 3), source, r, where, str(label_text).strip(), note))

    found = {}
    for cnd in sorted(cands, key=lambda x: (-x.conf, x.row)):
        found.setdefault(cnd.fid, cnd)
    return found


# ── Horizontal table layout (field names across a header row) ────────────────
def _extract_columns(book, name, warnings):
    head = book.rows_values(name, COLUMN_HEADER_SCAN_ROWS, COLUMN_MAX_COLS)
    if not head:
        return {}
    best_row, best_map = None, {}
    for r, row in enumerate(head):
        texts = [(c, v) for c, v in enumerate(row) if isinstance(v, str) and len(v.strip()) >= 2 and _to_num(v) is None]
        if len(texts) < 3:
            continue
        cmap = {}                                   # col -> (fid, conf, source, text)
        clf, thr = _get_classifier(warnings)
        pending = []
        for c, t in texts:
            pick = _model_pick([t], clf, thr)
            if pick:
                cmap[c] = (pick[0], pick[1], 'is_ml', t)
            else:
                pending.append((c, t))
        for c, t in pending:                        # safety net: rules
            m = _match_label(t)
            if m:
                cmap[c] = (m[0], m[1], 'is_column', t)
        # one column per field: keep the most confident
        by_fid = {}
        for c, (fid, conf, src, t) in cmap.items():
            if fid not in by_fid or conf > by_fid[fid][1]:
                by_fid[fid] = (c, conf, src, t)
        if len(by_fid) > len(best_map):
            best_row, best_map = r, by_fid
    if best_row is None or len(best_map) < 3:
        return {}

    body = book.rows_values(name, best_row + 1 + COLUMN_MAX_ROWS, COLUMN_MAX_COLS)[best_row + 1:]
    found = {}
    for fid, (c, conf, src, header) in best_map.items():
        nums = [n for n in (_to_num(row[c]) if c < len(row) else None for row in body) if n is not None]
        if not nums:
            continue
        avg = sum(nums) / len(nums)
        where = f"'{name}'!{_col_letter(c)}{best_row + 2}" + (f'…{_col_letter(c)}{best_row + 1 + len(body)}' if len(nums) > 1 else '')
        fin = _finalise_value(fid, avg, [header], where, warnings)
        if fin is None:
            continue
        value, note = fin
        note = (note + '; ' if note else '') + (f'mean of {len(nums)} readings' if len(nums) > 1 else '') or None
        found[fid] = Candidate(fid, value, round(conf, 3), src, best_row, where, str(header).strip(), note or None)
    return found if len(found) >= 3 else {}


def _sheet_names_standard(book, name):
    """True if the top of the sheet names the standard (e.g. '… (IS - 13980)')."""
    try:
        for row in book.rows_values(name, MARKER_SCAN_ROWS, VERTICAL_MAX_COLS):
            for v in row:
                if isinstance(v, str) and STANDARD_MARKER_RE.search(v):
                    return True
    except Exception:
        pass
    return False


# ── Public API ────────────────────────────────────────────────────────────────
def extract_is13890(file_bytes):
    """
    Given the raw bytes of an uploaded workbook, return:
        { 'extracted':     {field_id: value, ...},
          'missingFields': [{'id':..., 'label':...}, ...],      # required inputs not found
          'sheetName':     matching sheet name(s) or None,
          'fieldDetail':   {field_id: {label, header, source, confidence, sheet, cell, note}},
          'legend':        None   (kept for response compatibility; colours are not used),
          'warnings':      [str, ...] }
    Never raises for a workbook that simply holds no IS-13890 data — it returns an
    empty extraction. Only a genuinely unreadable file raises (the caller wraps
    this in its own try/except so it can never break the other detectors).

    Every sheet is scanned; the sheet that yields the most inputs is the primary
    one. Fields still missing are then filled from other sheets, but only from
    sheets that themselves look like an IS-13890 form (>= 4 fields).
    """
    empty = {'extracted': {}, 'missingFields': [{'id': f, 'label': FIELDS[f]['label']} for f in REQUIRED_IDS],
             'sheetName': None, 'fieldDetail': {}, 'legend': None, 'warnings': []}
    book = _Book(file_bytes)
    warnings = []
    results = []                              # (name, found{fid: Candidate}, names_standard)

    for name in book.names:
        try:
            found = _extract_vertical(book, name, warnings)
            if len(found) < MIN_FIELDS_TO_ACCEPT_SHEET:
                col_found = _extract_columns(book, name, warnings)
                if len(col_found) > len(found):
                    found = col_found
            if found:
                results.append((name, found, _sheet_names_standard(book, name)))
        except Exception as e:                # one odd sheet must not sink the others
            warnings.append(f"Sheet '{name}' skipped ({type(e).__name__}: {e}).")

    # GATE: a sheet that is not recognisably an IS-13890 form is dropped.
    gated = []
    for name, found, marked in results:
        distinctive = len(DISTINCTIVE_IDS & set(found))
        if marked or distinctive >= MIN_DISTINCTIVE_FIELDS:
            gated.append((name, found))
    if not gated:
        empty['warnings'] = warnings
        return empty

    gated.sort(key=lambda t: (len(t[1]), sum(c.conf for c in t[1].values())), reverse=True)
    primary_name, primary = gated[0]
    merged = {fid: (primary_name, c) for fid, c in primary.items()}
    used_sheets = [primary_name]
    for name, found in gated[1:]:
        if len(found) < MIN_FIELDS_TO_ACCEPT_SHEET:
            continue
        added = False
        for fid, c in found.items():
            if fid not in merged:
                merged[fid] = (name, c)
                added = True
        if added:
            used_sheets.append(name)

    if len(merged) < MIN_FIELDS_TO_ACCEPT_SHEET:     # too thin to trust — likely a coincidental match
        empty['warnings'] = warnings
        return empty

    extracted, detail = {}, {}
    for fid in FIELD_ORDER:
        if fid not in merged:
            continue
        sheet, c = merged[fid]
        extracted[fid] = c.value
        detail[fid] = {
            'label': FIELDS[fid]['label'],
            'header': f"{c.label} [{c.cell.split('!')[-1]}]",
            'source': c.source,
            'confidence': c.conf,
            'sheet': sheet,
            'cell': c.cell,
            'note': c.unit_note,
        }
        if c.unit_note:
            warnings.append(f"{FIELDS[fid]['label']} at {c.cell}: {c.unit_note}.")

    return {
        'extracted': extracted,
        'missingFields': [{'id': f, 'label': FIELDS[f]['label']} for f in REQUIRED_IDS if f not in extracted],
        'sheetName': ', '.join(dict.fromkeys(used_sheets)),
        'fieldDetail': detail,
        'legend': None,
        'warnings': warnings,
    }


if __name__ == '__main__':          # python -m routes.upload_is13890 path/to/book.xlsx
    import sys
    with open(sys.argv[1], 'rb') as fh:
        res = extract_is13890(fh.read())
    print(f"sheet: {res['sheetName']}")
    for fid, d in res['fieldDetail'].items():
        print(f"  {fid:4s} {res['extracted'][fid]:>14.6g}   {d['source']:10s} {d['confidence']:.2f}  {d['header']}")
    print('missing:', [m['id'] for m in res['missingFields']])
    for w in res['warnings']:
        print('warning:', w)
